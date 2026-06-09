#!/usr/bin/env python3
"""Generate formatted Word project report (Quiz Automation Platform stack)."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend" / ".docx_deps"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches


def add_title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("QUIZ AUTOMATION PLATFORM (QuizAI)")
    r.bold = True
    r.font.size = Pt(22)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(
        "An AI-Assisted Smart Quiz & Learning Analytics Web Application\n"
        "(FastAPI · React · SQLAlchemy · Google Gemini)"
    )
    r.font.size = Pt(14)
    for line in [
        "",
        "MASTER’S OF COMPUTER APPLICATIONS",
        "PANIPAT INSTITUTE OF ENGINEERING & TECHNOLOGY, PANIPAT",
        "BATCH (2024–26)",
        "",
        "SUBMITTED BY:",
        "[Student Name]",
        "Roll No.: [Roll Number]",
        "",
        "SUBMITTED TO:",
        "[Guide / HOD Name]",
        "Head of Department",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def body(doc, text):
    doc.add_paragraph(text)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def main():
    out = ROOT / "Quiz_Automation_Platform_Project_Report.docx"
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    add_title_page(doc)

    h1(doc, "Chapter 1 — Introduction")
    h2(doc, "1.1 Introduction")
    body(
        doc,
        "Educational institutions and training programs increasingly depend on online multiple-choice "
        "assessments for screening, internal exams, and skill evaluation. Paper-based or fragile quiz "
        "scripts fail under live-event load and create heavy manual work for faculty. The Quiz Automation "
        "Platform is a full-stack web application that lets teachers create and publish timed MCQ quizzes, "
        "lets students attempt assessments through a dedicated portal, and gives administrators control "
        "over platform-wide behaviour. Optional Google Gemini integration assists with question generation. "
        "The backend is built with FastAPI and SQLAlchemy; the frontend uses React (Vite) and Tailwind CSS. "
        "For development, data is stored in SQLite; for larger deployments, PostgreSQL with connection "
        "pooling is supported.",
    )
    h2(doc, "1.2 Project Overview")
    body(doc, "The system is divided into three primary roles:")
    bullets(
        doc,
        [
            "Administrator — platform settings (student analytics, retakes, rank tiers), aggregate analytics.",
            "Teacher — quiz lifecycle, question bank (manual, Excel/CSV, AI), publish, cohort analytics, bulk student import.",
            "Student — registration with profile fields, timed attempts, results, leaderboard, optional analytics.",
        ],
    )
    h2(doc, "1.3 Existing System Limitations")
    bullets(
        doc,
        [
            "Manual question entry only; slow for large banks.",
            "No unified AI-assisted authoring in legacy college scripts.",
            "Single login screen mixing staff and students causes role confusion.",
            "Instant answer reveal during exams reduces integrity.",
            "SQLite-only or un-pooled databases struggle under hundreds of simultaneous submits.",
        ],
    )
    h2(doc, "1.4 Proposed System")
    body(
        doc,
        "The proposed platform uses a decoupled SPA + REST API architecture. JWT authentication with "
        "bcrypt password hashing secures access. Teachers ingest questions manually, via spreadsheet, or "
        "through Gemini. Students use a separate login portal from staff. Administrators can disable "
        "student analytics and instant per-question feedback for exam-style attempts. Anti-cheat tab-switch "
        "logging warns candidates and can auto-submit after a threshold. Connection pooling and PostgreSQL "
        "are recommended when targeting live events with many concurrent users.",
    )
    h2(doc, "1.5 Problem Statement")
    bullets(
        doc,
        [
            "Reduce faculty time to build and publish MCQs at scale.",
            "Support large student cohorts with stable submission under load.",
            "Store student identity data (email, mobile, registration ID) for institutional records.",
            "Enforce role-appropriate login and configurable exam feedback modes.",
        ],
    )
    h2(doc, "1.6 Objectives")
    bullets(
        doc,
        [
            "Triple-path question creation: manual, bulk import, AI (topic/text/PDF).",
            "Timed quiz attempts with auto-submit on timeout.",
            "Teacher analytics and admin platform controls.",
            "Bulk student onboarding via Excel/CSV for live events.",
            "Secure authentication and API-first design for future scaling.",
        ],
    )
    h2(doc, "1.7 Scope")
    body(
        doc,
        "In scope: web-based MCQ quizzes, three roles, analytics dashboards, leaderboard, Gemini-assisted "
        "generation, Docker-based deployment option. Out of scope in current version: webcam AI proctoring, "
        "SMS notifications, PDF certificates, native mobile apps, and email OTP password recovery.",
    )
    h2(doc, "1.8 SDLC Model")
    body(doc, "The project follows an Agile incremental model: requirements, design, development, testing, deployment, maintenance.")

    doc.add_page_break()
    h1(doc, "Chapter 2 — Software Requirement Specification")
    h2(doc, "2.1 Introduction")
    body(doc, "This SRS describes functional and non-functional requirements for the Quiz Automation Platform.")
    h2(doc, "2.2 Users and Roles")
    bullets(
        doc,
        [
            "Admin — settings, platform analytics.",
            "Teacher — quizzes, questions, student import, quiz/cohort analytics.",
            "Student — register, attempt published quizzes, view results and leaderboard.",
        ],
    )
    h2(doc, "2.3 Core Features")
    bullets(
        doc,
        [
            "Authentication: register, login with portal (student vs staff), JWT sessions.",
            "Quiz management: create, edit, publish, question pool modes (all, first N, last N, random N).",
            "Question bank: CRUD, approval, bulk import, AI generate/explain/PDF.",
            "Attempts: start, save answers, submit, results, history, anti-cheat events.",
            "Analytics: student (toggle), teacher per-quiz and insights, admin overview.",
            "Platform settings: student analytics, retakes, rank tier cutoffs.",
        ],
    )
    h2(doc, "2.4 Non-Functional Requirements")
    bullets(
        doc,
        [
            "Security: bcrypt passwords, JWT, CORS configuration, secrets in .env.",
            "Performance: pooled DB connections; PostgreSQL for production load.",
            "Usability: responsive dashboards, light/dark theme support.",
            "Maintainability: modular FastAPI routers, React component structure.",
        ],
    )

    doc.add_page_break()
    h1(doc, "Chapter 3 — Technology Stack")
    h2(doc, "3.1 Frontend")
    bullets(
        doc,
        [
            "React 19 with Vite build tool.",
            "Tailwind CSS for styling; React Router for navigation.",
            "Native fetch API client with JWT in Authorization header.",
            "Framer Motion, Recharts, react-hot-toast for UX.",
        ],
    )
    h2(doc, "3.2 Backend")
    bullets(
        doc,
        [
            "Python 3.x, FastAPI, Uvicorn ASGI server.",
            "SQLAlchemy ORM; Pydantic schemas.",
            "python-jose (JWT), passlib/bcrypt.",
            "pandas + openpyxl for spreadsheet import.",
            "google-genai for Gemini API integration.",
        ],
    )
    h2(doc, "3.3 Database")
    bullets(
        doc,
        [
            "SQLite file (quiz_platform.db) for local development — no separate DB account required.",
            "PostgreSQL + psycopg2 recommended for production / ~1500-user events.",
            "Lightweight SQLite migrations in migrations_sqlite.py for schema upgrades.",
        ],
    )
    h2(doc, "3.4 Deployment")
    bullets(
        doc,
        [
            "Local: npm run dev (concurrent API + frontend).",
            "Optional: Docker Compose (Postgres, backend, frontend).",
            "Production: build frontend (VITE_API_URL), run Uvicorn with multiple workers behind nginx.",
        ],
    )

    doc.add_page_break()
    h1(doc, "Chapter 4 — System Design")
    h2(doc, "4.1 Architecture")
    body(
        doc,
        "Three-tier: React SPA → REST JSON API (FastAPI) → relational database. Gemini is called only "
        "from the backend so API keys are not exposed to browsers.",
    )
    h2(doc, "4.2 Data Flow (Summary)")
    body(
        doc,
        "Teacher creates quiz and questions → publishes → student starts attempt → answers saved via "
        "POST /attempts/save-answer → submit calculates score → result and analytics endpoints serve "
        "aggregated data. Public settings endpoint exposes non-secret flags (e.g. student analytics enabled).",
    )
    h2(doc, "4.3 Logical Database Entities")
    bullets(
        doc,
        [
            "users — roles, credentials, student profile, participant_code, mobile_digits.",
            "quizzes, questions — MCQ content, pool settings, publish flag.",
            "quiz_attempts, attempt_answers — scores, timestamps, selected options.",
            "platform_settings — singleton toggles for analytics, retakes, rank tiers.",
            "anticheat_logs — tab-switch and related events.",
        ],
    )
    h2(doc, "4.4 Diagrams")
    body(
        doc,
        "[Insert DFD Level 0/1, ER diagram, use case, sequence, and activity diagrams here — "
        "Draw.io / Lucidchart. Figures referenced in viva should match this implementation.]",
    )

    doc.add_page_break()
    h1(doc, "Chapter 5 — Design, Development & Testing")
    h2(doc, "5.1 Module Summary")
    bullets(
        doc,
        [
            "Auth module — /api/auth (register, login, me).",
            "Users module — profile, leaderboard, bulk-import-students.",
            "Quizzes & Questions — CRUD, publish, bulk-import.",
            "Attempts — start, questions, save-answer, submit, result, anticheat.",
            "Analytics — student, teacher, admin routes.",
            "AI — generate questions, explain mistake, PDF/content ingestion.",
            "Admin panel — settings patch, platform overview.",
        ],
    )
    h2(doc, "5.2 Testing Performed")
    bullets(
        doc,
        [
            "Manual end-to-end: teacher login → create quiz → add/import questions → publish → student attempt → submit → result.",
            "API smoke tests on health, auth, quizzes, analytics endpoints.",
            "Optional load_test.py for concurrent student simulation (tune for target load).",
        ],
    )
    h2(doc, "5.3 Repository")
    body(doc, "Source code: https://github.com/sneh-2create/quiz-automation--web.git")

    doc.add_page_break()
    h1(doc, "Chapter 6 — Conclusion & Future Scope")
    h2(doc, "6.1 Conclusion")
    body(
        doc,
        "The Quiz Automation Platform delivers a practical MCQ assessment workflow for teachers and students "
        "with AI-assisted authoring and analytics. It uses a modern Python/React stack suitable for MCA-level "
        "implementation and can be extended to PostgreSQL for institutional live events.",
    )
    h2(doc, "6.2 Future Enhancements")
    bullets(
        doc,
        [
            "Admin reset of a single student attempt (disaster recovery).",
            "Leaderboard and report export to Excel.",
            "Email OTP / password recovery.",
            "Webcam or AI-based proctoring.",
            "SMS/email result notifications and PDF certificates.",
            "Alembic migrations for long-lived PostgreSQL deployments.",
            "Formal load test report for 1500+ concurrent submits.",
        ],
    )
    h2(doc, "6.3 Bibliography")
    bullets(
        doc,
        [
            "FastAPI — https://fastapi.tiangolo.com",
            "SQLAlchemy — https://docs.sqlalchemy.org",
            "React — https://react.dev",
            "Google Gemini API — Google AI developer documentation",
            "Project repository (self-developed).",
        ],
    )

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
